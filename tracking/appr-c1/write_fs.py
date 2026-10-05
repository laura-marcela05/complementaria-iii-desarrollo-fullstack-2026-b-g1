# -*- coding: utf-8 -*-
"""Escribe F-S!P (appr). Por defecto SIMULA; --write para escribir."""
import openpyxl, json, io, sys, unicodedata, shutil, datetime

XLSX = r"C:\www\code-corhuila\repos_2026-b\activity-management\list-data\calification-term-1.xlsx"
SHEET = "F-S"
DO = "--write" in sys.argv
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

def norm(s):
    s = unicodedata.normalize('NFD', str(s))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return ' '.join(s.upper().split())

a = json.load(io.open("assessment_fs.json", encoding="utf-8"))
COMP = {int(k): v for k, v in a["completos"].items()}
INC = {int(k): v for k, v in a["incompletos"].items()}
SIN = {int(k): v for k, v in a["sin_fork"].items()}
roster = {r["row"]: r for r in json.load(io.open("roster_fs.json", encoding="utf-8"))}

wb = openpyxl.load_workbook(XLSX)
ws = wb[SHEET]
hdr = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
print("Encabezado en vivo:", {v: k for k, v in sorted(hdr.items(), key=lambda x: x[1])})
if hdr.get("appr") != 16:
    sys.exit(f"ABORTA: 'appr' no esta en P(16), esta en {hdr.get('appr')}.")
print("OK: F-S columna P = 'appr' confirmada en vivo.\n")

plan, errores = [], []
for row in sorted(COMP):   # SOLO quienes completaron las 5 semanas
    nom = ws.cell(row, 3).value
    if norm(nom) != norm(roster[row]["nombre"]):
        errores.append(f"fila {row}: nombre en vivo '{nom}' != '{roster[row]['nombre']}'"); continue
    hits = [r for r in range(2, ws.max_row + 1) if norm(ws.cell(r, 3).value or "") == norm(nom)]
    if len(hits) != 1:
        errores.append(f"fila {row}: el nombre cruza con {len(hits)} filas"); continue
    actual = ws.cell(row, 16).value
    if actual is not None:
        errores.append(f"fila {row}: ya tiene {actual!r}, no se sobrescribe"); continue
    nuevo = 0.6
    user = COMP[row][0]
    sem = "[1,2,3,4,5]"
    plan.append((row, nom.strip(), user, sem, nuevo))

print(f"{'fil':<5}{'estudiante':<38}{'usuario':<32}{'semanas':<16}{'P (appr)'}")
print("-" * 100)
for row, nom, user, sem, v in plan:
    print(f"{row:<5}{nom:<38}{user:<32}{sem:<16}{v}")
print(f"\nCeldas a escribir: {len(plan)}  (0.6: {sum(1 for p in plan if p[4]==0.6)} | 0: {sum(1 for p in plan if p[4]==0)})")
if errores:
    print("\n!!! BLOQUEOS:")
    for e in errores: print("   -", e)
print("\nSin fork — se dejan en blanco:")
for r in sorted(SIN):
    print(f"   fila {r:<3} {roster[r]['nombre']:<38} {SIN[r]}")

if not DO:
    print("\n=== SIMULACION. Archivo sin modificar. ===")
    sys.exit(0)
if errores:
    sys.exit("\nABORTA: hay bloqueos; no se escribe nada.")

stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
bak = XLSX.replace(".xlsx", f".backup-{stamp}.xlsx")
shutil.copy2(XLSX, bak)
print("\nRespaldo:", bak)
for row, nom, user, sem, v in plan:
    ws.cell(row, 16).value = v
try:
    wb.save(XLSX)
except PermissionError:
    sys.exit("ABORTA: el archivo esta abierto en Excel. Cierralo y vuelve a correr.")
print(f"Escritas {len(plan)} celdas en {SHEET}!P")

wb2 = openpyxl.load_workbook(XLSX); w2 = wb2[SHEET]
mal = [r for r, _, _, _, v in plan if w2.cell(r, 16).value != v]
frm = [r for r in range(2, w2.max_row + 1) if w2.cell(r, 3).value
       and not str(w2.cell(r, 17).value or "").startswith("=I")]
cd = wb2["C-D"]
cd_ok = sum(1 for r in range(2, 42) if cd.cell(r, 16).value == 0.6)
print(f"Verificacion -> mal escritas: {mal or 'ninguna'} | formulas Q rotas: {frm or 'ninguna'} | C-D intacta (0.6={cd_ok})")
