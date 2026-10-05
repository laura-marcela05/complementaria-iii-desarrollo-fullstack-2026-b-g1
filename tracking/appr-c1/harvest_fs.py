# -*- coding: utf-8 -*-
"""Cosecha completa de un curso: tree-diff (path+sha, todas las ramas) + CONFIG + extraccion de binarios."""
import subprocess, json, os, base64, io, re, sys, tempfile

SLUG = "complementaria-iii-desarrollo-fullstack-2026-b-g1"
BASE = f"code-corhuila/{SLUG}"
EV = rf"C:\www\code-corhuila\repos_2026-b\{SLUG}\tracking\appr-c1\evidencia"
DOCENTE = "ariel5253"
os.makedirs(EV, exist_ok=True)

def ghjson(path):
    p = subprocess.run(["gh", "api", path], capture_output=True)
    out = p.stdout.decode("utf-8", "replace").strip()
    if p.returncode != 0 or not out:
        return None, (p.stderr.decode("utf-8", "replace").strip() or "empty")
    try:
        return json.loads(out), None
    except Exception as ex:
        return None, f"parse:{ex}"

bt, err = ghjson(f"repos/{BASE}/git/trees/HEAD?recursive=1")
if err: sys.exit("no base tree: " + err)
BASE_SHA = {t["path"]: t["sha"] for t in bt["tree"] if t["type"] == "blob"}
print(f"esqueleto base: {len(BASE_SHA)} archivos", flush=True)

fk, err = ghjson(f"repos/{BASE}/forks?per_page=100")
if err: sys.exit("no forks: " + err)
forks = [{"user": f["owner"]["login"], "repo": f["full_name"], "default": f["default_branch"],
          "pushed": f["pushed_at"]} for f in fk if f["owner"]["login"] != DOCENTE]
print(f"forks de estudiantes: {len(forks)}", flush=True)

TEXT = {".md", ".txt", ".py", ".ipynb", ".csv", ".json", ".yml", ".yaml", ".html", ".css", ".js",
        ".jsx", ".ts", ".tsx", ".sql", ".java", ".php", ".vue", ".env", ".sh", ".xml", ".toml", ""}
BIN = {".pdf", ".docx", ".xlsx", ".xlsm", ".svg"}
SKIP_DIR = re.compile(r'(^|/)(node_modules|\.venv|venv|dist|build|\.next|vendor|__pycache__|\.git)/')
tmp = os.path.join(tempfile.gettempdir(), "fs_extract.bin")

def extract(raw, ext):
    try:
        if ext == ".pdf":
            import pdfplumber
            open(tmp, "wb").write(raw)
            with pdfplumber.open(tmp) as pdf:
                return "\n".join(f"--- pag {i+1} ---\n" + (p.extract_text() or "")
                                 for i, p in enumerate(pdf.pages[:40]))
        if ext == ".docx":
            import docx
            open(tmp, "wb").write(raw)
            d = docx.Document(tmp)
            out = [p.text for p in d.paragraphs]
            for t in d.tables:
                for r in t.rows: out.append(" | ".join(c.text for c in r.cells))
            return "\n".join(out)
        if ext in (".xlsx", ".xlsm"):
            import openpyxl
            open(tmp, "wb").write(raw)
            wb = openpyxl.load_workbook(tmp, data_only=True)
            out = []
            for ws in wb.worksheets:
                out.append(f"=== hoja {ws.title} ===")
                for row in ws.iter_rows(max_row=80, values_only=True):
                    c = [str(x) for x in row if x is not None]
                    if c: out.append(" | ".join(c))
            return "\n".join(out)
        if ext == ".svg":
            return " ".join(re.findall(r'>([^<>]{2,})<', raw.decode("utf-8", "replace")))
    except Exception as ex:
        return f"[[EXTRACT ERROR {ext}: {ex}]]"
    return None

results = []
for i, f in enumerate(forks, 1):
    user, repo = f["user"], f["repo"]
    rec = {**f, "config": None, "full_name_cfg": None, "added": [], "branches": [], "errors": [], "skipped": 0}

    cfg, cerr = ghjson(f"repos/{user}/{user}/contents/README.md")
    if cfg and "content" in cfg:
        try:
            txt = base64.b64decode(cfg["content"]).decode("utf-8", "replace")
            rec["config"] = txt
            m = re.search(r'FULL[_ ]?NAME\s*\**\s*[:|=]\s*\**\s*(.+)', txt, re.I)
            if m:
                rec["full_name_cfg"] = m.group(1).strip().replace("*", "").replace("`", "").replace("#", "").strip()
        except Exception as ex:
            rec["errors"].append(f"config: {ex}")
    else:
        rec["errors"].append(f"sin repo de perfil/README: {cerr}")

    br, berr = ghjson(f"repos/{repo}/branches?per_page=100")
    branches = [b["name"] for b in br] if br else [f["default"]]
    if berr: rec["errors"].append(f"branches: {berr}")
    rec["branches"] = branches

    seen = {}
    for b in branches:
        tr, terr = ghjson(f"repos/{repo}/git/trees/{b}?recursive=1")
        if terr:
            rec["errors"].append(f"tree {b}: {terr}"); continue
        for t in tr.get("tree", []):
            if t["type"] != "blob": continue
            p = t["path"]
            if p.endswith(".gitkeep"): continue
            if p in BASE_SHA and BASE_SHA[p] == t["sha"]: continue   # identico al base
            if SKIP_DIR.search(p):
                rec["skipped"] += 1; continue                        # dependencias, no es entrega
            if t["sha"] in seen:
                if b not in seen[t["sha"]]["branches"]: seen[t["sha"]]["branches"].append(b)
                continue
            seen[t["sha"]] = {"path": p, "sha": t["sha"], "size": t.get("size", 0), "branches": [b],
                              "modified_base": p in BASE_SHA}
    rec["added"] = list(seen.values())

    sdir = os.path.join(EV, user)
    for a in rec["added"]:
        ext = os.path.splitext(a["path"])[1].lower()
        a["ext"] = ext
        if a["size"] > 900_000: continue
        if ext not in TEXT and ext not in BIN: continue
        blob, oerr = ghjson(f"repos/{repo}/git/blobs/{a['sha']}")
        if oerr:
            a["error"] = oerr; continue
        try:
            raw = base64.b64decode(blob["content"])
        except Exception as ex:
            a["error"] = str(ex); continue
        txt = extract(raw, ext) if ext in BIN else raw.decode("utf-8", "replace")
        if txt is None: continue
        a["text"] = txt
        os.makedirs(sdir, exist_ok=True)
        safe = a["path"].replace("/", "__").replace("\\", "__")
        if ext in BIN: safe += ".txt"
        try:
            io.open(os.path.join(sdir, safe), "w", encoding="utf-8").write(txt)
        except OSError as ex:
            a["error"] = f"save: {ex}"
    results.append(rec)
    print(f"[{i:>2}/{len(forks)}] {user:<32} +{len(rec['added']):<4} ramas:{len(branches)} "
          f"omit:{rec['skipped']:<5} {'ERR' if rec['errors'] else ''}", flush=True)

json.dump(results, io.open("harvest_fs.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("\nOK -> harvest_fs.json ; evidencia en", EV)
